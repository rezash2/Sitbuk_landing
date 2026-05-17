from django.contrib import admin

from .forms import BlogAdminForm
from .models import (
    BaleBotScenario,
    BaleOperatorReplyTemplate,
    BlogPost,
    ContactMessage,
    DemoRequest,
    FAQItem,
    HomeContentItem,
    HomeHeroContent,
    LeadRequest,
    LeadFollowUpActivity,
    MediaAsset,
    NewsletterSubscription,
    PageContent,
    PageContentItem,
    PageBuilderSection,
    PricingComparisonRow,
    PricingPlan,
    SiteSettings,
    SiteRedirect,
)


def _export_rows_as_excel(modeladmin, request, queryset, filename, columns):
    """Export admin rows as XLSX when openpyxl exists; otherwise return UTF-8 CSV safely."""
    try:
        from openpyxl import Workbook
        from openpyxl.utils import get_column_letter
    except Exception:
        import csv
        from django.http import HttpResponse

        response = HttpResponse(content_type='text/csv; charset=utf-8-sig')
        response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
        response.write('\ufeff')
        writer = csv.writer(response)
        writer.writerow([label for label, _ in columns])
        for obj in queryset:
            writer.writerow([getter(obj) for _, getter in columns])
        return response

    from django.http import HttpResponse

    wb = Workbook()
    ws = wb.active
    ws.title = 'Leads'
    ws.append([label for label, _ in columns])
    for obj in queryset:
        ws.append([getter(obj) for _, getter in columns])
    for col in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 24
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    wb.save(response)
    return response


def export_leads(modeladmin, request, queryset):
    columns = [
        ('نام', lambda o: o.full_name),
        ('تلفن', lambda o: o.phone),
        ('شرکت', lambda o: o.company),
        ('ایمیل', lambda o: o.email),
        ('وضعیت', lambda o: o.get_status_display()),
        ('اولویت', lambda o: o.get_priority_display()),
        ('مسئول پیگیری', lambda o: o.assigned_to),
        ('صفحه مبدا', lambda o: o.source_page),
        ('UTM Source', lambda o: o.utm_source),
        ('UTM Campaign', lambda o: o.utm_campaign),
        ('زمان ثبت', lambda o: o.created_at.strftime('%Y-%m-%d %H:%M') if o.created_at else ''),
        ('زمان پیگیری بعدی', lambda o: o.follow_up_at.strftime('%Y-%m-%d %H:%M') if o.follow_up_at else ''),
        ('آخرین تماس', lambda o: o.last_contacted_at.strftime('%Y-%m-%d %H:%M') if o.last_contacted_at else ''),
        ('توضیحات', lambda o: o.note),
        ('یادداشت داخلی', lambda o: o.internal_note),
    ]
    return _export_rows_as_excel(modeladmin, request, queryset, 'sitbuk-leads', columns)
export_leads.short_description = 'خروجی Excel از لیدهای انتخاب‌شده'


def mark_leads_contacted(modeladmin, request, queryset):
    queryset.update(status=LeadRequest.STATUS_CONTACTED)
mark_leads_contacted.short_description = 'تغییر وضعیت به پیگیری شده'


def mark_leads_qualified(modeladmin, request, queryset):
    queryset.update(status=LeadRequest.STATUS_QUALIFIED)
mark_leads_qualified.short_description = 'تغییر وضعیت به واجد شرایط'


def mark_leads_won(modeladmin, request, queryset):
    queryset.update(status=LeadRequest.STATUS_WON)
mark_leads_won.short_description = 'تغییر وضعیت به تبدیل شده'


def mark_leads_lost(modeladmin, request, queryset):
    queryset.update(status=LeadRequest.STATUS_LOST)
mark_leads_lost.short_description = 'تغییر وضعیت به رد شده'


def export_demo_requests(modeladmin, request, queryset):
    columns = [
        ('نام', lambda o: o.full_name),
        ('تلفن', lambda o: o.phone),
        ('ایمیل', lambda o: o.email),
        ('شرکت', lambda o: o.company),
        ('نوع دمو', lambda o: o.get_demo_type_display()),
        ('وضعیت', lambda o: o.get_status_display()),
        ('اولویت', lambda o: o.get_priority_display()),
        ('لینک دمو', lambda o: o.demo_access_url),
        ('اعتبار لینک', lambda o: o.demo_access_expires_at.strftime('%Y-%m-%d %H:%M') if getattr(o, 'demo_access_expires_at', None) else ''),
        ('تعداد ورود', lambda o: getattr(o, 'demo_launch_count', 0)),
        ('آخرین دمو', lambda o: getattr(o, 'last_demo_target', '')),
        ('صفحه مبدا', lambda o: o.source_page),
        ('زمان ثبت', lambda o: o.created_at.strftime('%Y-%m-%d %H:%M') if o.created_at else ''),
        ('توضیحات', lambda o: o.note),
        ('یادداشت داخلی', lambda o: o.internal_note),
    ]
    return _export_rows_as_excel(modeladmin, request, queryset, 'sitbuk-demo-requests', columns)
export_demo_requests.short_description = 'خروجی Excel از درخواست‌های دمو'


def mark_demo_reviewing(modeladmin, request, queryset):
    queryset.update(status=DemoRequest.STATUS_REVIEWING)
mark_demo_reviewing.short_description = 'تغییر وضعیت به در حال بررسی'


def mark_demo_link_ready(modeladmin, request, queryset):
    queryset.update(status=DemoRequest.STATUS_LINK_READY)
mark_demo_link_ready.short_description = 'تغییر وضعیت به لینک آماده'


@admin.register(DemoRequest)
class DemoRequestAdmin(admin.ModelAdmin):
    list_display = ('full_name', 'phone', 'email', 'company', 'demo_type', 'status', 'priority', 'source_page', 'created_at')
    list_editable = ('status', 'priority')
    search_fields = ('full_name', 'phone', 'email', 'company', 'note', 'internal_note', 'demo_access_token', 'demo_access_url')
    list_filter = ('demo_type', 'status', 'priority', 'source_page', 'created_at')
    date_hierarchy = 'created_at'
    readonly_fields = ('demo_access_token', 'demo_launch_count', 'last_demo_target', 'page_url', 'referrer', 'utm_source', 'utm_campaign', 'ip_address', 'user_agent', 'created_at', 'updated_at')
    actions = [export_demo_requests, mark_demo_reviewing, mark_demo_link_ready]
    fieldsets = (
        ('اطلاعات درخواست دمو', {
            'fields': ('full_name', 'phone', 'email', 'company', 'demo_type', 'note')
        }),
        ('پیگیری و لینک دمو', {
            'fields': ('status', 'priority', 'assigned_to', 'internal_note', 'demo_access_token', 'demo_access_url', 'demo_access_expires_at', 'demo_launch_count', 'last_demo_target', 'demo_link_sent_at', 'demo_entered_at')
        }),
        ('ردیابی منبع', {
            'classes': ('collapse',),
            'fields': ('source_page', 'page_url', 'referrer', 'utm_source', 'utm_campaign', 'ip_address', 'user_agent')
        }),
        ('زمان‌ها', {
            'classes': ('collapse',),
            'fields': ('created_at', 'updated_at')
        }),
    )


@admin.register(LeadRequest)
class LeadRequestAdmin(admin.ModelAdmin):
    list_display = ('full_name', 'phone', 'company', 'status', 'priority', 'assigned_to', 'source_page', 'follow_up_at', 'created_at')
    list_editable = ('status', 'priority', 'assigned_to')
    search_fields = ('full_name', 'phone', 'company', 'email', 'note', 'internal_note', 'utm_source', 'utm_campaign')
    list_filter = ('status', 'priority', 'source_page', 'created_at', 'follow_up_at')
    date_hierarchy = 'created_at'
    readonly_fields = ('page_url', 'referrer', 'utm_source', 'utm_campaign', 'ip_address', 'user_agent', 'created_at', 'updated_at')
    actions = [export_leads, mark_leads_contacted, mark_leads_qualified, mark_leads_won, mark_leads_lost]
    fieldsets = (
        ('اطلاعات مخاطب', {
            'fields': ('full_name', 'phone', 'company', 'email', 'note')
        }),
        ('پیگیری فروش', {
            'fields': ('status', 'priority', 'assigned_to', 'follow_up_at', 'last_contacted_at', 'internal_note')
        }),
        ('ردیابی منبع', {
            'classes': ('collapse',),
            'fields': ('source_page', 'page_url', 'referrer', 'utm_source', 'utm_campaign', 'ip_address', 'user_agent')
        }),
        ('زمان‌ها', {
            'classes': ('collapse',),
            'fields': ('created_at', 'updated_at')
        }),
    )


@admin.register(LeadFollowUpActivity)
class LeadFollowUpActivityAdmin(admin.ModelAdmin):
    list_display = ('lead', 'activity_type', 'result', 'user', 'next_follow_up_at', 'created_at')
    search_fields = ('lead__full_name', 'lead__phone', 'lead__company', 'note')
    list_filter = ('activity_type', 'result', 'created_at', 'next_follow_up_at')
    date_hierarchy = 'created_at'
    readonly_fields = ('created_at',)



def export_contact_messages(modeladmin, request, queryset):
    columns = [
        ('نام', lambda o: o.full_name),
        ('تلفن', lambda o: o.phone),
        ('ایمیل', lambda o: o.email),
        ('شرکت', lambda o: o.company),
        ('موضوع', lambda o: o.subject),
        ('وضعیت', lambda o: o.get_status_display()),
        ('اولویت', lambda o: o.get_priority_display()),
        ('صفحه مبدا', lambda o: o.source_page),
        ('زمان ثبت', lambda o: o.created_at.strftime('%Y-%m-%d %H:%M') if o.created_at else ''),
        ('پیام', lambda o: o.message),
        ('یادداشت داخلی', lambda o: o.internal_note),
    ]
    return _export_rows_as_excel(modeladmin, request, queryset, 'sitbuk-contact-messages', columns)
export_contact_messages.short_description = 'خروجی Excel از پیام‌های انتخاب‌شده'


def mark_messages_replied(modeladmin, request, queryset):
    queryset.update(status=ContactMessage.STATUS_REPLIED)
mark_messages_replied.short_description = 'تغییر وضعیت به پاسخ داده شده'


def mark_messages_closed(modeladmin, request, queryset):
    queryset.update(status=ContactMessage.STATUS_CLOSED)
mark_messages_closed.short_description = 'تغییر وضعیت به بسته شده'


@admin.register(NewsletterSubscription)
class NewsletterSubscriptionAdmin(admin.ModelAdmin):
    list_display = ('email', 'created_at')
    search_fields = ('email',)
    list_filter = ('created_at',)
    date_hierarchy = 'created_at'


@admin.register(ContactMessage)
class ContactMessageAdmin(admin.ModelAdmin):
    list_display = ('full_name', 'subject', 'phone', 'email', 'status', 'priority', 'source_page', 'created_at')
    list_editable = ('status', 'priority')
    search_fields = ('full_name', 'subject', 'phone', 'email', 'company', 'message', 'internal_note')
    list_filter = ('status', 'priority', 'source_page', 'created_at')
    date_hierarchy = 'created_at'
    readonly_fields = ('page_url', 'referrer', 'ip_address', 'user_agent', 'created_at', 'updated_at')
    actions = [export_contact_messages, mark_messages_replied, mark_messages_closed]
    fieldsets = (
        ('اطلاعات پیام', {
            'fields': ('full_name', 'phone', 'email', 'company', 'subject', 'message')
        }),
        ('پیگیری', {
            'fields': ('status', 'priority', 'internal_note')
        }),
        ('ردیابی منبع', {
            'classes': ('collapse',),
            'fields': ('source_page', 'page_url', 'referrer', 'ip_address', 'user_agent')
        }),
        ('زمان‌ها', {
            'classes': ('collapse',),
            'fields': ('created_at', 'updated_at')
        }),
    )


@admin.register(BlogPost)
class BlogPostAdmin(admin.ModelAdmin):
    form = BlogAdminForm
    list_display = ('title', 'category', 'is_featured', 'is_published', 'robots', 'published_at')
    list_filter = ('category', 'is_featured', 'is_published', 'published_at')
    search_fields = ('title', 'summary', 'content', 'seo_title', 'seo_description', 'seo_keywords')
    prepopulated_fields = {'slug': ('title',)}
    ordering = ('-published_at',)
    fieldsets = (
        ('محتوا', {
            'fields': ('title', 'slug', 'category', 'summary', 'content', 'reading_time', 'accent', 'is_featured', 'is_published', 'published_at')
        }),
        ('SEO و اشتراک‌گذاری', {
            'classes': ('collapse',),
            'fields': ('seo_title', 'seo_description', 'seo_keywords', 'og_image', 'canonical_url', 'robots')
        }),
    )


@admin.register(FAQItem)
class FAQItemAdmin(admin.ModelAdmin):
    list_display = ('question', 'sort_order', 'is_active')
    list_filter = ('is_active',)
    search_fields = ('question', 'answer')
    ordering = ('sort_order', 'id')


@admin.register(SiteSettings)
class SiteSettingsAdmin(admin.ModelAdmin):
    list_display = ('site_name', 'support_phone', 'sales_phone', 'support_email', 'updated_at')
    fieldsets = (
        ('اطلاعات عمومی', {
            'fields': ('site_name', 'support_phone', 'sales_phone', 'support_email', 'address', 'working_hours', 'footer_about')
        }),
        ('شبکه‌های اجتماعی', {
            'fields': ('whatsapp_number', 'telegram_url', 'instagram_url', 'linkedin_url')
        }),
        ('SEO عمومی سایت', {
            'classes': ('collapse',),
            'fields': ('seo_title_suffix', 'default_meta_description', 'default_meta_keywords', 'default_og_image', 'default_og_image_alt', 'robots_policy')
        }),
    )




@admin.register(SiteRedirect)
class SiteRedirectAdmin(admin.ModelAdmin):
    list_display = ('source_path', 'target_url', 'status_code', 'is_active', 'hit_count', 'last_used_at', 'updated_at')
    list_editable = ('status_code', 'is_active')
    search_fields = ('source_path', 'target_url', 'internal_note')
    list_filter = ('status_code', 'is_active', 'updated_at')
    readonly_fields = ('hit_count', 'last_used_at', 'created_at', 'updated_at')

@admin.register(HomeHeroContent)
class HomeHeroContentAdmin(admin.ModelAdmin):
    list_display = ('title_highlight', 'kicker_primary', 'primary_button_label', 'is_active', 'updated_at')
    list_filter = ('is_active',)
    fieldsets = (
        ('تیتر و پیام اصلی', {
            'fields': ('is_active', 'eyebrow', 'kicker_primary', 'kicker_secondary', 'title_prefix', 'title_highlight', 'title_suffix', 'description')
        }),
        ('دکمه‌ها و تصویر', {
            'fields': ('primary_button_label', 'primary_button_url', 'secondary_button_label', 'secondary_button_url_name', 'hero_image')
        }),
    )


@admin.register(HomeContentItem)
class HomeContentItemAdmin(admin.ModelAdmin):
    list_display = ('title', 'section', 'value', 'badge', 'icon', 'sort_order', 'is_active')
    list_filter = ('section', 'is_active')
    search_fields = ('title', 'subtitle', 'description', 'value', 'badge', 'icon')
    ordering = ('section', 'sort_order', 'id')
    list_editable = ('sort_order', 'is_active')

@admin.register(PageContent)
class PageContentAdmin(admin.ModelAdmin):
    list_display = ('page_key', 'page_title', 'robots', 'schema_type', 'is_active', 'updated_at')
    list_filter = ('page_key', 'is_active')
    search_fields = ('page_title', 'page_description', 'seo_title', 'seo_description', 'seo_keywords')
    fieldsets = (
        ('صفحه', {
            'fields': ('page_key', 'is_active', 'page_kicker', 'page_title', 'page_description')
        }),
        ('SEO و اشتراک‌گذاری', {
            'fields': ('seo_title', 'seo_description', 'seo_keywords', 'canonical_path', 'robots', 'og_type', 'schema_type', 'og_image', 'og_image_alt')
        }),
        ('تصویر Hero', {
            'fields': ('hero_image', 'hero_alt')
        }),
    )


@admin.register(PageContentItem)
class PageContentItemAdmin(admin.ModelAdmin):
    list_display = ('title', 'page_key', 'section', 'value', 'badge', 'icon', 'sort_order', 'is_active')
    list_filter = ('page_key', 'section', 'is_active')
    search_fields = ('title', 'subtitle', 'description', 'value', 'badge', 'icon', 'image', 'url')
    ordering = ('page_key', 'section', 'sort_order', 'id')
    list_editable = ('sort_order', 'is_active')
    fieldsets = (
        ('محل نمایش', {
            'fields': ('page_key', 'section', 'sort_order', 'is_active')
        }),
        ('محتوا', {
            'fields': ('title', 'subtitle', 'description', 'value', 'badge', 'icon', 'image', 'url')
        }),
    )





@admin.register(PageBuilderSection)
class PageBuilderSectionAdmin(admin.ModelAdmin):
    list_display = ('page_key', 'section_key', 'title', 'layout', 'sort_order', 'is_active', 'is_published', 'updated_at')
    list_editable = ('title', 'layout', 'sort_order', 'is_active', 'is_published')
    search_fields = ('page_key', 'section_key', 'title', 'description')
    list_filter = ('page_key', 'layout', 'is_active', 'is_published')


@admin.register(PricingPlan)
class PricingPlanAdmin(admin.ModelAdmin):
    list_display = ('name', 'context', 'monthly_price', 'annual_price', 'accent', 'sort_order', 'is_active', 'updated_at')
    list_filter = ('context', 'is_active', 'accent')
    search_fields = ('name', 'subtitle', 'tag', 'description', 'features_text')
    ordering = ('context', 'sort_order', 'id')
    list_editable = ('sort_order', 'is_active')


@admin.register(PricingComparisonRow)
class PricingComparisonRowAdmin(admin.ModelAdmin):
    list_display = ('label', 'table_key', 'group_title', 'sort_order', 'is_active', 'updated_at')
    list_filter = ('table_key', 'group_title', 'is_active')
    search_fields = ('label', 'group_title', 'value_1', 'value_2', 'value_3', 'value_4', 'value_5', 'value_6')
    ordering = ('table_key', 'group_title', 'sort_order', 'id')
    list_editable = ('sort_order', 'is_active')


@admin.register(MediaAsset)
class MediaAssetAdmin(admin.ModelAdmin):
    list_display = ('title', 'asset_type', 'usage_key', 'is_active', 'created_at')
    list_filter = ('asset_type', 'is_active', 'created_at')
    search_fields = ('title', 'alt_text', 'usage_key', 'description', 'file')
    readonly_fields = ('created_at', 'updated_at')



from .models import DashboardAuditLog


@admin.register(DashboardAuditLog)
class DashboardAuditLogAdmin(admin.ModelAdmin):
    list_display = ('created_at', 'username', 'section', 'action', 'object_repr', 'ip_address')
    list_filter = ('section', 'action', 'created_at')
    search_fields = ('username', 'section', 'object_repr', 'path', 'ip_address', 'user_agent')
    readonly_fields = ('user', 'username', 'action', 'section', 'object_repr', 'path', 'method', 'ip_address', 'user_agent', 'metadata', 'created_at')


@admin.register(BaleBotScenario)
class BaleBotScenarioAdmin(admin.ModelAdmin):
    list_display = ('title', 'key', 'action', 'match_mode', 'sort_order', 'is_active', 'updated_at')
    list_editable = ('sort_order', 'is_active')
    search_fields = ('title', 'key', 'trigger_keywords', 'response_text')
    list_filter = ('action', 'match_mode', 'is_active')


@admin.register(BaleOperatorReplyTemplate)
class BaleOperatorReplyTemplateAdmin(admin.ModelAdmin):
    list_display = ('title', 'category', 'sort_order', 'is_active', 'updated_at')
    list_editable = ('sort_order', 'is_active')
    search_fields = ('title', 'text')
    list_filter = ('category', 'is_active')
